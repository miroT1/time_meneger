from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT_DIR = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT_DIR / "data" / "time_tracking.xlsx"

SHEETS = {
    "Projects": ["ID", "Название", "Категория", "Статус", "Планируемое время (ч)", "Описание", "Создан"],
    "Tasks": ["ID", "ID проекта", "Название", "Статус", "Описание", "Создан"],
    "Sessions": ["ID", "Дата", "ID проекта", "ID задачи", "Время начала", "Время окончания", "Длительность (ч)", "Что сделано", "Источник"],
    "Statistics": ["Показатель", "Значение"],
}


def ensure_workbook():
    """Создаёт необходимые листы, сохраняя все существующие листы и данные."""
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        wb = load_workbook(WORKBOOK_PATH)
    else:
        wb = Workbook()
        if wb.active:
            wb.remove(wb.active)

    for sheet_name, headers in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
            _format_header(ws)
        else:
            ws = wb[sheet_name]
            if ws.max_row == 1 and all(ws.cell(1, col).value is None for col in range(1, ws.max_column + 1)):
                ws.delete_rows(1, ws.max_row)
                ws.append(headers)
                _format_header(ws)
            elif ws.max_row == 1 and ws.cell(1, 1).value is None:
                ws.append(headers)
                _format_header(ws)

    wb.save(WORKBOOK_PATH)
    wb.close()


def _format_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A81")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        max_len = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 35)


def read_table(sheet_name):
    ensure_workbook()
    wb = load_workbook(WORKBOOK_PATH, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return __import__("pandas").DataFrame()
    headers = list(rows[0])
    data = [row for row in rows[1:] if any(value is not None for value in row)]
    return __import__("pandas").DataFrame(data, columns=headers)


def _next_id(sheet_name, prefix):
    wb = load_workbook(WORKBOOK_PATH, data_only=True)
    ws = wb[sheet_name]
    ids = [str(ws.cell(row, 1).value) for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value]
    wb.close()
    numbers = []
    for value in ids:
        if value.startswith(prefix):
            try:
                numbers.append(int(value[len(prefix):]))
            except ValueError:
                pass
    return f"{prefix}{max(numbers, default=0) + 1:04d}"


def _append(sheet_name, values):
    wb = load_workbook(WORKBOOK_PATH)
    ws = wb[sheet_name]
    ws.append(values)
    _format_header(ws)
    wb.save(WORKBOOK_PATH)
    wb.close()


def add_project(name, category, status, target_hours, description):
    existing = read_table("Projects")
    if not existing.empty and existing["Название"].astype(str).str.casefold().eq(name.casefold()).any():
        raise ValueError("Проект с таким названием уже существует.")
    project_id = _next_id("Projects", "P")
    _append("Projects", [project_id, name, category, status, float(target_hours), description, datetime.now().isoformat(timespec="seconds")])
    return project_id


def add_task(project_id, name, status, description):
    projects = read_table("Projects")
    if projects.empty or project_id not in projects["ID"].astype(str).tolist():
        raise ValueError("Выбранный проект не найден.")
    tasks = read_table("Tasks")
    if not tasks.empty:
        duplicate = (
            tasks["ID проекта"].astype(str).eq(project_id)
            & tasks["Название"].astype(str).str.casefold().eq(name.casefold())
        ).any()
        if duplicate:
            raise ValueError("В этом проекте уже есть задача с таким названием.")
    task_id = _next_id("Tasks", "T")
    _append("Tasks", [task_id, project_id, name, status, description, datetime.now().isoformat(timespec="seconds")])
    return task_id


def add_session(session_date, project_id, task_id, start_iso, end_iso, notes, source="manual"):
    projects = read_table("Projects")
    tasks = read_table("Tasks")
    if projects.empty or project_id not in projects["ID"].astype(str).tolist():
        raise ValueError("Проект не найден.")
    if tasks.empty:
        raise ValueError("Задача не найдена.")
    task_rows = tasks[tasks["ID"].astype(str).eq(task_id)]
    if task_rows.empty or str(task_rows.iloc[0]["ID проекта"]) != project_id:
        raise ValueError("Задача не принадлежит выбранному проекту.")
    start = datetime.fromisoformat(start_iso)
    end = datetime.fromisoformat(end_iso)
    duration = (end - start).total_seconds() / 3600
    if duration <= 0:
        raise ValueError("Длительность должна быть положительной.")
    session_id = _next_id("Sessions", "S")
    _append("Sessions", [session_id, session_date, project_id, task_id, start_iso, end_iso, round(duration, 4), notes, source])
    return session_id


def list_projects():
    df = read_table("Projects")
    if df.empty:
        return []
    return df.fillna("").to_dict(orient="records")


def list_tasks(project_id=None):
    df = read_table("Tasks")
    if df.empty:
        return []
    if project_id is not None:
        df = df[df["ID проекта"].astype(str).eq(project_id)]
    return df.fillna("").to_dict(orient="records")
